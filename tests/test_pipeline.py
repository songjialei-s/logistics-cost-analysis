import json
import shutil
import sqlite3
import tempfile
import unittest
from contextlib import closing
from pathlib import Path
import pandas as pd
from src.etl import ROOT, run, normalize
from src.generate_sample import generate

class PipelineTest(unittest.TestCase):
    def test_invalid_fields_and_fx(self):
        cfg=json.loads((ROOT/'config/providers.json').read_text())[0]
        fx=json.loads((ROOT/'config/fx_rates.json').read_text())
        base=dict(zip(cfg['mapping'].values(),['DEMO-X','2026-01-01','1,234.50','1.2','2','DELIVERY']))
        record=normalize(base,cfg,fx,2)
        self.assertEqual((record[5],record[7]),(123450,123450))
        for key,value in [('ship_date','2025-01-01'),('amount','-2'),('weight','0'),('item_count','1.5'),('category','UNKNOWN')]:
            row=base.copy(); row[cfg['mapping'][key]]=value
            with self.assertRaises((ValueError,KeyError)): normalize(row,cfg,fx,2)

    def test_conflict_rolls_back(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp); shutil.copytree(ROOT/'config',root/'config'); generate(root)
            file=root/'data/sample/AsterParcel.xlsx'
            wb=load_workbook(file); ws=wb.active
            row=[cell.value for cell in ws[2]]; row[2]='999'; ws.append(row); wb.save(file)
            with self.assertRaisesRegex(ValueError,'Conflicting duplicate'): run(root)
            with closing(sqlite3.connect(root/'runtime/demo.sqlite')) as con:
                self.assertEqual(con.execute('SELECT COUNT(*) FROM charges').fetchone()[0],0)
                self.assertEqual(con.execute('SELECT COUNT(*) FROM processed_files').fetchone()[0],0)

    def test_end_to_end(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp); shutil.copytree(ROOT/'config',root/'config'); generate(root)
            first=run(root)
            self.assertEqual((first['inserted'],first['duplicates'],first['rejected']),(144,6,18))
            self.assertEqual((first['total_parcels'],first['anomaly_lines']),(72,1))
            second=run(root); self.assertEqual((second['files_skipped'],second['inserted']),(6,0))
            with closing(sqlite3.connect(root/'runtime/demo.sqlite')) as con:
                self.assertEqual(con.execute('SELECT SUM(amount_cny_minor) FROM charges').fetchone()[0],794700)
                for provider in json.loads((root/'config/providers.json').read_text()):
                    weight=con.execute('SELECT weight_kg FROM charges WHERE provider=? AND tracking_id LIKE ? LIMIT 1',(provider['provider'],'%-1-1')).fetchone()[0]
                    self.assertAlmostEqual(weight,1.2,places=3)
                sql='\n'.join(line for line in (ROOT/'sql/analysis.sql').read_text().splitlines() if not line.lstrip().startswith('--'))
                queries=[s.strip() for s in sql.split(';') if s.strip()]
                results=[con.execute(q).fetchall() for q in queries]
                self.assertEqual(len(results),5)
                self.assertEqual([r[1] for r in results[0]],[2360,2488,3099])
                self.assertIsNone(results[0][0][2])
                self.assertAlmostEqual(results[0][1][2],128/2360)
                self.assertAlmostEqual(sum(r[1] for r in results[2]),1)
                self.assertAlmostEqual(results[3][0][0],99.7083333333)
                self.assertAlmostEqual(results[3][0][1],49.8541666667)
                self.assertEqual(len(results[4]),1)
            self.assertEqual(len(pd.read_csv(root/'data/output/charges.csv')),144)
            self.assertEqual(len(pd.read_csv(root/'data/output/parcels.csv')),72)
            # Changed files rejected before database changes.
            file=root/'data/sample/AsterParcel.xlsx'; file.write_bytes(file.read_bytes()+b' ')
            with self.assertRaisesRegex(ValueError,'changed'): run(root)
            self.assertEqual(run(root,rebuild=True)['inserted'],144)

if __name__=='__main__': unittest.main()
